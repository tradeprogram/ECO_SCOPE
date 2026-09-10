"""에코스코프 가이드북 — 배경지식이 없는 독자를 위한 해설서.

    python scripts/build_guidebook.py [출력폴더]

서식은 `하수범_공모전/물길잡이_충청남도/물길잡이_가이드북.pdf` 를 따릅니다.
    A4 · 맑은 고딕 · 흑백 · 표는 얇은 실선 2열(용어 | 뜻)
    표지 → 목차 → 제1장 용어사전 → 본문 → 부록 A/B
    별표(★)는 이 시스템을 이해하는 데 꼭 필요한 장에 붙입니다.

수치는 web/data/*.json 과 소스의 상수에서 직접 읽습니다. 손으로 옮겨 적으면
판독을 다시 돌렸을 때 문서만 옛 숫자로 남습니다.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document                                   # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT              # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK     # noqa: E402
from docx.oxml import OxmlElement                           # noqa: E402
from docx.oxml.ns import qn                                 # noqa: E402
from docx.shared import Cm, Pt, RGBColor                    # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

FONT = "맑은 고딕"
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)


# ── 수치 ────────────────────────────────────────────────────────────────────

def numbers() -> dict:
    d = lambda p: json.loads((ROOT / "web/data" / p).read_text(encoding="utf-8"))
    s, v = d("summary.json"), d("validation.json")
    ys = d("wetland_years.json")
    grades = d("ecomap_grades.json")
    orbits = d("orbit_history.json")
    from nie.wetlands import aggregate as A
    from nie.rs import s1_water as W

    cov = sum(y.get("n_covered") or 0 for y in ys)
    obs = sum(y["n_obs"] for y in ys)
    area = v["area_level_vs_variation"]
    return {
        "s": s, "v": v, "area": area,
        "cov": cov, "obs": obs, "cov_pct": round(cov / obs * 100, 1),
        "n_grades": len(grades), "n_orbits": len(orbits),
        "COVER_RATIO": A.COVER_RATIO, "VV_VEG_DB": A.VV_VEG_DB,
        "MIN_OBS": A.MIN_OBS_FOR_YEAR, "MIN_BASE": A.MIN_BASELINE_RATIO,
        "MIN_OPEN_HA": A.MIN_OPEN_HA, "CONF": A.CONF_DISAGREE_FRAC,
        "VV_MAX": W.WATER_VV_DB_MAX, "PIX": W.S1_PIXEL_M,
        "OTSU_RANGE": W.OTSU_VALID_RANGE,
        "open": v["by_state"]["open"],
    }


# ── 서식 도우미 ─────────────────────────────────────────────────────────────

def setup(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.4)
    sec.left_margin = sec.right_margin = Cm(2.5)

    # 쪽번호 — 아래 가운데
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); r.font.name = FONT; r.font.size = Pt(9); r.font.color.rgb = GREY
    for el, attr in (("w:fldChar", {"w:fldCharType": "begin"}),
                     ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
        e = OxmlElement(el)
        if el == "w:instrText":
            e.set(qn("xml:space"), "preserve"); e.text = " PAGE "
        for k, val in (attr or {}).items():
            e.set(qn(k), val)
        r._r.append(e)


def para(doc, text="", *, size=10, bold=False, align=None, before=0, after=0,
         color=None, indent=0, spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = spacing
    if indent:
        pf.left_indent = Pt(indent)
    if text:
        r = p.add_run(text)
        r.font.name = FONT; r.font.size = Pt(size); r.bold = bold
        r.font.color.rgb = color or INK
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p


def chapter(doc, text):
    para(doc, text, size=16, bold=True, before=22, after=10)


def section(doc, text):
    para(doc, text, size=11.5, bold=True, before=14, after=5)


def body(doc, text):
    para(doc, text, after=6)


def _cell(c, text, *, bold=False, size=9.5):
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    r.font.name = FONT; r.font.size = Pt(size); r.bold = bold
    r.font.color.rgb = INK
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def term_table(doc, rows, head=("용어", "뜻"), widths=(3.6, 12.4)):
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell(t.rows[0].cells[0], head[0], bold=True)
    _cell(t.rows[0].cells[1], head[1], bold=True)
    for a, b in rows:
        r = t.add_row()
        _cell(r.cells[0], a)
        _cell(r.cells[1], b)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def figure(doc, path: Path, caption: str, width_cm=15.5):
    if not path.exists():
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    para(doc, caption, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
         after=10, color=GREY, spacing=1.3)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
